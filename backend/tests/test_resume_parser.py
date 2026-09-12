import os
import sys
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def test_resume_parser_extracts_text():
    from app.services.resume_parser import ResumeParser

    resume_text = """
    John Doe
    john.doe@gmail.com
    SKILLS
    Python, FastAPI, PostgreSQL, Docker
    EXPERIENCE
    Senior Developer at TechNova (2019-2024)
    Backend Developer at QuantumSoft (2016-2019)
    EDUCATION
    B.Tech in Computer Science
    CERTIFICATIONS
    AWS Certified Solutions Architect
    """
    parsed = ResumeParser.extract_candidate_info(resume_text)
    assert "python" in parsed["skills"].lower()
    assert parsed["years_of_experience"] >= 4
    assert "b.tech" in parsed["education"].lower()
    assert "aws" in parsed["certifications"].lower()


def test_resume_parser_extracts_from_docx(tmp_path):
    from docx import Document
    from app.services.resume_parser import ResumeParser

    file_path = tmp_path / "resume.docx"
    doc = Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("jane.smith@gmail.com")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, React, SQL")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Software Engineer at SaaSify (2021-2024)")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("B.Tech in Computer Science")
    doc.save(str(file_path))

    text = ResumeParser.extract_from_docx(str(file_path))
    assert text is not None
    assert "python" in text.lower()

    info = ResumeParser.extract_candidate_info(text)
    assert info["years_of_experience"] >= 1